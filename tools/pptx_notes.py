#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx"]
# ///
r"""把 pptx 备注栏的讲稿换成 speech.md 的原文。

pdf2pptx.py 的备注是从 dual.pdf 右半提取的，而 PDF 里存的是排版后的结果：
每个视觉行是一个独立 block，提取出来就成了句子中间的硬断行（行与行之间
还被 "\n\n" 连成不同段落）；em dash 前后的空格是排版间距不是空格字符，
提取后变成 "succeeded —but the"。这些都没法从 PDF 侧可靠还原。

但讲稿原文就在 speech.md，所以不必还原——直接拿原文替换。页和讲稿块的
对应关系由备注栏顶上的面包屑给出（note 4/16 · step 2/4），那是 beamer
排版时写进去的，不用猜。

面包屑所在的那一行（连同它前面的标题）保留，只换正文。面包屑之后本来就没有
正文的页（比如结束页，面包屑停在上一帧的值）保持原样，不会被填进重复的讲稿。

用法：python3 pptx_notes.py <pptx> <speech.md>
"""
import io
import re
import sys

from pptx import Presentation

# 面包屑由 preamble.tex 的 note page 模板排出：<crumbnote> <段>/<总段数>
# [· <crumbstep> <块>/<总块数>]。中英文两套用词都认。
# 讲稿分块时排成「动画 2/3」，没分块但按过空格时只排「动画 2」——后者的数字是
# overlay 序号，不是块号，这一帧整段讲稿就一块，所以那种情况下块号取 1。
# 面包屑和帧标题排在同一行（标题 \hfill 面包屑），pdf2pptx.py 把它们并成一行文本，
# 所以只锚行尾：面包屑前面可以有标题，也可以什么都没有。
BREADCRUMB = re.compile(
    r"(?:^|\s)(?:note|讲稿)\s*(\d+)/\d+"
    r"(?:\s*·\s*(?:step|动画)\s*(\d+)(?:/(\d+))?)?\s*$")


def load_blocks(path):
    """speech.md → {段号: [块文本, ...]}，切法与 speech2tex.py 一致。"""
    raw = io.open(path, encoding="utf-8").read()
    parts = re.split(r"^\*\*P(\d+)\.\*\*\s*", raw, flags=re.M)
    if len(parts) < 3:
        sys.exit(f"{path} 里没找到 **Pn.** 段落标记")

    out = {}
    for num, body in zip(parts[1::2], parts[2::2]):
        blocks = []
        for chunk in re.split(r"\[>\]", body.strip()):
            if not chunk.strip():
                continue
            # 段内换行只是源文件的折行 → 合成空格；空行是真的分段 → 留空行
            paras = [" ".join(p.split())
                     for p in re.split(r"\n\s*\n", chunk) if p.strip()]
            blocks.append("\n\n".join(paras))
        out[int(num)] = blocks
    return out


def split_head(text):
    """(标题+面包屑, 面包屑后是否还有正文, 段号, 块号)；没有面包屑则 None。"""
    lines = text.splitlines()
    for i, ln in enumerate(lines):
        m = BREADCRUMB.search(ln.strip())
        if m:
            para = int(m.group(1))
            # 只有带分母的「块号/总块数」才是真的块号
            step = int(m.group(2)) if (m.group(2) and m.group(3)) else 1
            head = "\n".join(l for l in lines[:i + 1] if l.strip())
            has_body = any(l.strip() for l in lines[i + 1:])
            return head, has_body, para, step
    return None


def main():
    if len(sys.argv) < 3:
        sys.exit("用法：pptx_notes.py <pptx> <speech.md>")
    pptx_path, speech_path = sys.argv[1], sys.argv[2]

    blocks = load_blocks(speech_path)
    prs = Presentation(pptx_path)

    replaced = skipped = 0
    for n, slide in enumerate(prs.slides, 1):
        if not slide.has_notes_slide:
            continue
        frame = slide.notes_slide.notes_text_frame
        info = split_head(frame.text)
        if info is None:
            continue
        head, has_body, para, step = info
        if not has_body:          # 结束页这类：面包屑停在上一帧，本页本来无稿
            skipped += 1
            continue
        if para not in blocks or step > len(blocks[para]):
            sys.exit(f"第 {n} 页面包屑指向 P{para} 第 {step} 块，speech.md 里没有")
        frame.text = f"{head}\n\n{blocks[para][step - 1]}"
        replaced += 1

    # python-pptx 的默认模板会在文件属性里留下库作者的名字和 "generated
    # using python-pptx"——对着一份对外交付的 deck，这两项清掉更体面
    cp = prs.core_properties
    cp.last_modified_by = ""
    cp.comments = ""
    cp.revision = 1

    prs.save(pptx_path)
    print(f"  备注换成 speech.md 原文：{replaced} 页"
          + (f"，{skipped} 页无稿跳过" if skipped else ""))


if __name__ == "__main__":
    main()
